from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / 'CINEON_DEPLOY_A_Z.md'
OUTPUT = ROOT / 'CINEON_DEPLOY_A_Z.docx'
doc = Document()
# The runtime template carries a blue Title border. Remove inherited rules.
for style in doc.styles:
    for border in style._element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.2677), Inches(11.6929)
sec.top_margin = sec.bottom_margin = Inches(.72)
sec.left_margin = sec.right_margin = Inches(.72)
sec.header_distance = sec.footer_distance = Inches(.30)

for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3']:
    s = doc.styles[name]
    s.font.name = 'Arial'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Arial')
normal = doc.styles['Normal']
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.10
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.widow_control = True
doc.styles['Title'].font.size = Pt(23)
doc.styles['Title'].paragraph_format.space_after = Pt(12)
for name, size in [('Heading 1', 16), ('Heading 2', 12), ('Heading 3', 11)]:
    s = doc.styles[name]
    s.font.size = Pt(size)
    s.font.bold = True
    s.paragraph_format.space_before = Pt(13)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.keep_with_next = True

code = doc.styles.add_style('Command', 1)
code.font.name = 'Consolas'
code.font.size = Pt(8.5)
code.font.color.rgb = RGBColor(0, 0, 0)
code.paragraph_format.line_spacing = 1.03
code.paragraph_format.space_before = Pt(4)
code.paragraph_format.space_after = Pt(8)
code.paragraph_format.left_indent = Inches(.08)
code.paragraph_format.right_indent = Inches(.04)

doc.core_properties.title = 'Hướng dẫn triển khai và vận hành Cineon'
doc.core_properties.subject = 'Triển khai cineon.me trên Ubuntu 24.04 với VPS 1 vCPU 4 GB RAM'
doc.core_properties.author = 'Cineon'
doc.core_properties.keywords = 'Cineon, VPS, deployment, Docker, Caddy, Sentry, monitoring'

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Cineon   ').font.size = Pt(8)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)

def text_clean(s):
    return s.replace('**', '').replace('`', '')

def shade(element, color):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    element.append(shd)

def add_table(rows):
    headers = rows[0]
    n = len(headers)
    table = doc.add_table(rows=1, cols=n)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [2.45, 4.35] if n == 2 else [1.55, 1.72, 3.53]
    if n not in (2, 3): widths = [6.8 / n] * n
    borders = OxmlElement('w:tblBorders')
    for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        e = OxmlElement('w:' + edge)
        for k, v in [('val','single'),('sz','4'),('color','D9D9D9')]: e.set(qn('w:'+k), v)
        borders.append(e)
    table._tbl.tblPr.append(borders)
    for i, vals in enumerate(rows):
        row = table.rows[0] if i == 0 else table.add_row()
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement('w:cantSplit'); trPr.append(cant)
        if i == 0:
            repeat = OxmlElement('w:tblHeader'); trPr.append(repeat)
        for j, val in enumerate(vals):
            cell = row.cells[j]
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                e = OxmlElement('w:'+side); e.set(qn('w:w'), '90'); e.set(qn('w:type'), 'dxa'); margins.append(e)
            tcPr.append(margins)
            shade(tcPr, 'E8EDF2' if i == 0 else ('F7F9FB' if i % 2 == 0 else 'FFFFFF'))
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.line_spacing = 1.06
            p.paragraph_format.keep_with_next = i == 0
            r = p.add_run(text_clean(val))
            r.font.name = 'Arial'; r.font.size = Pt(9)
            r.font.bold = i == 0
            r.font.color.rgb = RGBColor(0,0,0)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

lines = SOURCE.read_text(encoding='utf-8').splitlines()
i = 0
while i < len(lines):
    line = lines[i]
    if not line.strip(): i += 1; continue
    if line.startswith('```'):
        block=[]; i += 1
        while i < len(lines) and not lines[i].startswith('```'):
            block.append(lines[i]); i += 1
        p = doc.add_paragraph('\n'.join(block), style='Command')
        p.paragraph_format.keep_together = len(block) <= 14
        shade(p._p.get_or_add_pPr(), 'F2F4F6')
    elif line.startswith('|'):
        rows=[]
        while i < len(lines) and lines[i].startswith('|'):
            cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+', c) for c in cells): rows.append(cells)
            i += 1
        add_table(rows); continue
    elif line.startswith('# '):
        doc.add_paragraph(line[2:], style='Title')
    elif line.startswith('### '):
        doc.add_paragraph(line[4:], style='Heading 2')
    elif line.startswith('## '):
        doc.add_paragraph(line[3:], style='Heading 1')
    elif line.startswith('- '):
        p=doc.add_paragraph(text_clean(line[2:]))
        p.paragraph_format.left_indent=Inches(.12)
    else:
        p=doc.add_paragraph(text_clean(line))
        if line.endswith(':'):
            p.paragraph_format.keep_with_next = True
    i += 1

doc.save(OUTPUT)
print(f'DOCX_CREATED {OUTPUT}')
print(f'PARAGRAPHS {len(doc.paragraphs)} TABLES {len(doc.tables)}')
